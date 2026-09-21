from models import Resultado
from fastapi import FastAPI, UploadFile, File, Depends, HTTPException
from sqlalchemy.orm import Session
from sqlalchemy import func
from openpyxl import load_workbook
from datetime import datetime
from io import BytesIO
from pydantic import BaseModel
import os
from fastapi.middleware.cors import CORSMiddleware
from scrapers.orquestador import sincronizar_desde_scraping, sincronizar_historico_desde_scraping
import models_miercoles  
import crud
import models_viernes
import models_scraping
import models_historico4                   
from routers import historico4 as router_historico4
from services.historico4 import backfill_desde_resultado, resumen_historico
from sqlalchemy import text
from database import SessionLocal
from services.analisis import (
    construir_analisis_completo,
    convertir_resultado,
    eliminar_resultado,
    importar_resultados_excel,
    importar_historico_excel,
    sincronizar_desde_api_datos_gov,
    construir_vista_excel,
    registrar_resultado_manual,
    obtener_resultados, 
    obtener_resultado_por_id,  
    reconstruir_estado_terminacion,
    ResultadoNoEncontrado,       # <- nuevo
    ResultadoDemasiadoAntiguo, 
)
  
from services import viernes as servicios_viernes
from services import miercoles as servicios_miercoles
from apscheduler.schedulers.background import BackgroundScheduler

from database import engine, Base

Base.metadata.create_all(bind=engine)

 

app = FastAPI()
 
frontend_url = os.getenv("FRONTEND_URL")
 
origenes_permitidos = [
    "http://localhost:5173",
    "http://127.0.0.1:5173",
]
if frontend_url:
    origenes_permitidos.append(frontend_url)
 
app.add_middleware(
    CORSMiddleware,
    allow_origins=origenes_permitidos,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.include_router(router_historico4.router)
# =========================================================
# CONEXIÓN A BASE DE DATOS
# =========================================================

def get_db():
    db = SessionLocal()

    try:
        yield db
    finally:
        db.close()

class ResultadoManualIn(BaseModel):
    fecha: str          # admite "YYYY-MM-DD" o "DD/MM/YYYY"
    numero: str
    loteria: str = "Manual"

@app.post("/debug/backfill-historico4")
def backfill_historico4_endpoint(db: Session = Depends(get_db)):
    agregados = backfill_desde_resultado(db)
    resumen = resumen_historico(db)
    return {
        "mensaje": f"{agregados} resultados agregados al histórico de 4 cifras",
        "resumen": resumen,
    }    


@app.post("/resultados/manual")
def crear_resultado_manual(
    datos: ResultadoManualIn,
    db: Session = Depends(get_db),
):
    try:
        resultado = registrar_resultado_manual(
            db, datos.fecha, datos.numero, datos.loteria
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    analisis = construir_analisis_completo(db)
    vista = construir_vista_excel(db)

    return {
        "mensaje": "Resultado agregado correctamente",
        "resultado": convertir_resultado(resultado),
        "analisis": list(analisis.values()),
        "vista_excel": vista,
    }


@app.post("/debug/reconstruir-estado")
def reconstruir_estado_endpoint(db: Session = Depends(get_db)):
    reconstruir_estado_terminacion(db)
    analisis = construir_analisis_completo(db)
    vista = construir_vista_excel(db)
    return {
        "mensaje": "Estado de terminaciones reconstruido correctamente",
        "analisis": list(analisis.values()),
        "vista_excel": vista,
    }

class ResultadoManualMiercolesIn(BaseModel):
    fecha: str          # admite "YYYY-MM-DD" o "DD/MM/YYYY"
    numero: str
    loteria: str        # "Meta" | "Valle" | "Manizales"


@app.post("/miercoles/manual")
def crear_resultado_manual_miercoles(
    datos: ResultadoManualMiercolesIn,
    db: Session = Depends(get_db),
):
    try:
        servicios_miercoles.registrar_resultado_manual_miercoles(
            db, datos.loteria, datos.fecha, datos.numero
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    return {
        "mensaje": "Resultado agregado correctamente",
        "vista": servicios_miercoles.construir_vista_miercoles(db),
    }


class ResultadoManualViernesIn(BaseModel):
    fecha: str          # admite "YYYY-MM-DD" o "DD/MM/YYYY"
    numero: str
    loteria: str 

    
# =========================================================
# INICIO
# =========================================================

@app.get("/")
def inicio():
    return {
        "mensaje": "Lottery Analyzer API funcionando 🚀"
    }


# =========================================================
# CREAR ARCHIVO
# =========================================================

@app.post("/archivos")
def crear_archivo(
    nombre: str,
    nombre_original: str,
    db: Session = Depends(get_db)
):
    ahora = datetime.now()

    archivo = crud.crear_archivo(
        db,
        nombre=nombre,
        nombre_original=nombre_original,
        fecha_creacion=ahora,
        fecha_actualizacion=ahora
    )

    return archivo


# =========================================================
# LISTAR ARCHIVOS
# =========================================================

@app.get("/archivos")
def listar_archivos(
    db: Session = Depends(get_db)
):
    return crud.obtener_archivos(db)

# =========================================================
# OBTENER RESULTADO POR ID
# =========================================================

@app.get("/resultados/{resultado_id}")
def obtener_resultado(resultado_id: int, db: Session = Depends(get_db)):
    try:
        resultado = obtener_resultado_por_id(db, resultado_id)
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))

    return convertir_resultado(resultado)


# =========================================================
# OBTENER ARCHIVO
# =========================================================

@app.get("/archivos/{archivo_id}")
def obtener_archivo(
    archivo_id: int,
    db: Session = Depends(get_db)
):
    archivo = crud.obtener_archivo(db, archivo_id)

    if not archivo:
        raise HTTPException(
            status_code=404,
            detail="Archivo no encontrado"
        )

    return archivo


# =========================================================
# ELIMINAR ARCHIVO
# =========================================================

@app.delete("/archivos/{archivo_id}")
def eliminar_archivo(
    archivo_id: int,
    db: Session = Depends(get_db)
):
    archivo = crud.eliminar_archivo(db, archivo_id)

    if not archivo:
        raise HTTPException(
            status_code=404,
            detail="Archivo no encontrado"
        )

    return {
        "mensaje": "Archivo eliminado correctamente"
    }

# =========================================================
# LISTAR RESULTADOS
# =========================================================

@app.get("/resultados")
def listar_resultados(
    terminacion: str | None = None,
    loteria: str | None = None,
    limite: int = 50,
    db: Session = Depends(get_db),
):
    resultados = obtener_resultados(db, terminacion=terminacion, loteria=loteria, limite=limite)
    return [convertir_resultado(r) for r in resultados]

# =========================================================
# ELIMINAR RESULTADO
# =========================================================

@app.delete("/resultados/{resultado_id}")
def borrar_resultado(resultado_id: int, db: Session = Depends(get_db)):
    try:
        eliminar_resultado(db, resultado_id)
    except ResultadoNoEncontrado as e:
        raise HTTPException(status_code=404, detail=str(e))
    except ResultadoDemasiadoAntiguo as e:
        raise HTTPException(status_code=403, detail=str(e))

    analisis = construir_analisis_completo(db)
    vista = construir_vista_excel(db)

    return {
        "mensaje": "Resultado eliminado correctamente",
        "analisis": list(analisis.values()),
        "vista_excel": vista,
    }

# =========================================================
# SUBIR Y LEER EXCEL (inspección rápida, sin insertar nada)
# =========================================================

@app.post("/excel/upload")
async def subir_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):

    # -----------------------------------------------------
    # VALIDAR EXTENSIÓN
    # -----------------------------------------------------

    if not file.filename.lower().endswith(
        (".xlsx", ".xlsm", ".xltx", ".xltm")
    ):
        raise HTTPException(
            status_code=400,
            detail="El archivo debe ser un Excel válido"
        )

    # -----------------------------------------------------
    # LEER ARCHIVO
    # -----------------------------------------------------

    contenido = await file.read()

    try:
        workbook = load_workbook(
            filename=BytesIO(contenido),
            read_only=True,
            data_only=True
        )

    except Exception:
        raise HTTPException(
            status_code=400,
            detail="No se pudo leer el archivo Excel"
        )

    # -----------------------------------------------------
    # INFORMACIÓN DE LAS HOJAS
    # -----------------------------------------------------

    hojas = []

    for worksheet in workbook.worksheets:

        primeras_filas = []

        for row in worksheet.iter_rows(
            min_row=1,
            max_row=5,
            values_only=True
        ):
            primeras_filas.append(list(row))

        hojas.append({
            "nombre": worksheet.title,
            "filas": worksheet.max_row,
            "columnas": worksheet.max_column,
            "primeras_filas": primeras_filas
        })

    workbook.close()

    # -----------------------------------------------------
    # CREAR REGISTRO DEL ARCHIVO
    # -----------------------------------------------------

    ahora = datetime.now()

    archivo = crud.crear_archivo(
        db,
        nombre=file.filename,
        nombre_original=file.filename,
        fecha_creacion=ahora,
        fecha_actualizacion=ahora
    )

    # -----------------------------------------------------
    # RESPUESTA
    # -----------------------------------------------------

    return {
        "mensaje": "Excel leído y guardado correctamente",
        "archivo_id": archivo.id,
        "nombre_archivo": file.filename,
        "hojas": hojas
    }


# =========================================================
# IMPORTAR RESULTADOS NUEVOS DESDE EXCEL (esto sí inserta y analiza)
# =========================================================

@app.post("/excel/importar-resultados")
async def importar_resultados(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    if not file.filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        raise HTTPException(
            status_code=400,
            detail="El archivo debe ser un Excel válido",
        )

    contenido = await file.read()
    ahora = datetime.now()

    # 1. Creamos el registro del archivo (archivo_id es obligatorio en Resultado)
    archivo = crud.crear_archivo(
        db,
        nombre=file.filename,
        nombre_original=file.filename,
        fecha_creacion=ahora,
        fecha_actualizacion=ahora,
    )

    # 2. Parseamos e insertamos los resultados
    try:
        nuevos = importar_resultados_excel(db, archivo.id, contenido)
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"No se pudo procesar el archivo: {e}",
        )

    # 3. Devolvemos el análisis ya recalculado, listo para el frontend
    analisis = construir_analisis_completo(db)

    return {
        "mensaje": f"{len(nuevos)} resultados importados correctamente",
        "archivo_id": archivo.id,
        "nuevos": [convertir_resultado(r) for r in nuevos],
        "analisis": list(analisis.values()),
    }


# =========================================================
# OBTENER EL ANÁLISIS COMPLETO (para pintar la tabla al abrir la app)
# =========================================================

@app.get("/analisis")
def obtener_analisis(db: Session = Depends(get_db)):
    analisis = construir_analisis_completo(db)
    return list(analisis.values())


scheduler = BackgroundScheduler()
 
def tarea_sincronizacion_scraping():
    db = SessionLocal()
    try:
        nuevos, reporte = sincronizar_desde_scraping(db)
        print(f"[scraping] {len(nuevos)} resultados nuevos — {reporte.resumen_texto()}")
    finally:
        db.close()

def tarea_sincronizacion_diaria():
    db = SessionLocal()
    try:
        nuevos = sincronizar_desde_api_datos_gov(db)
        print(f"[sincronización automática] {len(nuevos)} resultados nuevos")
    finally:
        db.close()
 
 
@app.on_event("startup")
def iniciar_scheduler():
    scheduler.add_job(tarea_sincronizacion_diaria, "cron", hour=7, minute=0)
    scheduler.add_job(tarea_sincronizacion_scraping, "cron", hour=7, minute=15)
    scheduler.start()

 
 
@app.on_event("shutdown")
def detener_scheduler():
    scheduler.shutdown()


@app.post("/scrapers/sincronizar")
def sincronizar_scraping_manual(db: Session = Depends(get_db)):
    nuevos, reporte = sincronizar_historico_desde_scraping(db, dias_atras=60)

    analisis = construir_analisis_completo(db)
    vista = construir_vista_excel(db)

    respuesta = {
        "mensaje": f"{len(nuevos)} resultados nuevos sincronizados por scraping",
        "loterias_ok": reporte.exitosos,
        "loterias_fallidas": reporte.fallidos,
        "analisis": list(analisis.values()),
        "vista_excel": vista,
    }

    if reporte.hubo_fallos:
        respuesta["advertencia"] = (
            f"No se pudo leer el resultado de: {', '.join(reporte.fallidos)}. "
            "Revisa esos scrapers (sitio caído o cambió el HTML)."
        )

    return respuesta
 

@app.get("/scrapers/estado")
def obtener_estado_scraping(db: Session = Depends(get_db)):
    fila = (
        db.query(models_scraping.EstadoUltimaCorridaScraping)
        .order_by(models_scraping.EstadoUltimaCorridaScraping.id.desc())
        .first()
    )
 
    if not fila:
        return {"hay_datos": False}
 
    return {
        "hay_datos": True,
        "fecha_ejecucion": fila.fecha_ejecucion.isoformat(),
        "loterias_ok": fila.loterias_ok.split(",") if fila.loterias_ok else [],
        "loterias_fallidas": fila.loterias_fallidas.split(",") if fila.loterias_fallidas else [],
    }
 
@app.post("/api-loterias/sincronizar")
def sincronizar_manual(db: Session = Depends(get_db)):
    """
    Dispara la sincronización manualmente (útil para probar sin
    esperar al cron de las 7am, o para un botón de 'actualizar ahora'
    en el frontend).
    """
    nuevos = sincronizar_desde_api_datos_gov(db)
    analisis = construir_analisis_completo(db)
 
    return {
        "mensaje": f"{len(nuevos)} resultados nuevos sincronizados",
        "analisis": list(analisis.values()),
    }
@app.get("/vista-excel")
def obtener_vista_excel(db: Session = Depends(get_db)):
    """Ventana 2: los 3 grupos de columnas + tabla amarilla, automáticos."""
    return construir_vista_excel(db)



@app.get("/debug/conteo-por-terminacion")
@app.get("/debug/fecha-maxima-global")
def debug_fecha_maxima_global(db: Session = Depends(get_db)):
    maxima = db.query(func.max(Resultado.fecha)).scalar()
    total = db.query(func.count(Resultado.id)).scalar()
    return {
        "fecha_maxima_en_resultado": maxima.isoformat() if maxima else None,
        "total_filas": total,
    }
def debug_conteo(db: Session = Depends(get_db)):
    conteos = (
        db.query(Resultado.numero, func.count(Resultado.id))
        .group_by(Resultado.numero)
        .all()
    )
    mapa = {n: c for n, c in conteos}
    faltantes = [f"{i:02d}" for i in range(100) if mapa.get(f"{i:02d}", 0) < 3]
    return {
        "conteos": {f"{i:02d}": mapa.get(f"{i:02d}", 0) for i in range(100)},
        "terminaciones_con_menos_de_3": faltantes,
        "total_faltantes": len(faltantes),
    }
 
@app.post("/excel/importar-historico")
async def importar_historico(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    if not file.filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        raise HTTPException(
            status_code=400,
            detail="El archivo debe ser un Excel válido",
        )
 
    contenido = await file.read()
    ahora = datetime.now()
 
    archivo = crud.crear_archivo(
        db,
        nombre=file.filename,
        nombre_original=file.filename,
        fecha_creacion=ahora,
        fecha_actualizacion=ahora,
    )
 
    try:
        nuevos = importar_historico_excel(db, archivo.id, contenido)
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"No se pudo procesar el archivo: {e}",
        )
 
    analisis = construir_analisis_completo(db)
 
    return {
        "mensaje": f"{len(nuevos)} resultados históricos importados correctamente",
        "archivo_id": archivo.id,
        "analisis": list(analisis.values()),
    }

@app.get("/debug/conexion")
def debug_conexion(db: Session = Depends(get_db)):
    fila = db.execute(text(
        "SELECT current_database(), current_schema(), inet_server_addr()::text, inet_server_port()"
    )).fetchone()
    columnas = db.execute(text(
        "SELECT column_name FROM information_schema.columns WHERE table_name = 'resultados_miercoles'"
    )).fetchall()
    return {
        "base_datos": fila[0],
        "esquema": fila[1],
        "host": fila[2],
        "puerto": fila[3],
        "columnas_resultados_miercoles": [c[0] for c in columnas],
    }
@app.post("/miercoles/importar-excel")
async def importar_excel_miercoles_endpoint(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    if not file.filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        raise HTTPException(status_code=400, detail="El archivo debe ser un Excel válido")

    contenido = await file.read()
    archivo = servicios_miercoles.crear_archivo_miercoles(db, file.filename)

    try:
        procesadas, diagnostico = servicios_miercoles.importar_excel_miercoles(db, archivo.id, contenido)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo procesar el archivo: {e}")

    # --- TEMPORAL: conteo por lotería para diagnóstico ---
    from collections import Counter
    conteo_por_loteria = Counter(loteria for loteria, _ in procesadas)

    return {
        "mensaje": f"{len(procesadas)} filas importadas, {len(diagnostico)} filas omitidas",
        "conteo_por_loteria": dict(conteo_por_loteria),  # <-- nuevo
        "archivo_id": archivo.id,
        "omitidas": diagnostico,  # <-- sin el [:20], quiero verlas TODAS
        "vista": servicios_miercoles.construir_vista_miercoles(db),
    }
@app.get("/miercoles/historial/{loteria}")
def obtener_historial_miercoles(loteria: str, db: Session = Depends(get_db)):
    resultados = (
        db.query(models_miercoles.ResultadoMiercoles)
        .filter(models_miercoles.ResultadoMiercoles.loteria == loteria)
        .order_by(models_miercoles.ResultadoMiercoles.id)   # <-- orden real del excel
        .all()
    )

    contador_por_numero = {}
    filas = []

    for r in resultados:
        contador_por_numero[r.numero] = contador_por_numero.get(r.numero, 0) + 1
        cantidad = r.cantidad if r.cantidad is not None else contador_por_numero[r.numero]

        filas.append({
            "id": r.id,
            "fecha": r.fecha.isoformat(),
            "numero": r.numero,
            "cantidad": cantidad,
            # solo se muestra si de verdad viene del excel (columnas N/O de Meta)
            "fecha_pen": r.fecha_pen.isoformat() if r.fecha_pen else None,
            "numero_pen": r.numero_pen,
        })

    return filas

@app.get("/miercoles/vista")
def obtener_vista_miercoles_endpoint(db: Session = Depends(get_db)):
    return servicios_miercoles.construir_vista_miercoles(db)


@app.post("/viernes/manual")
def crear_resultado_manual_viernes(
    datos: ResultadoManualViernesIn,
    db: Session = Depends(get_db),
):
    try:
        servicios_viernes.registrar_resultado_manual_viernes(
            db, datos.loteria, datos.fecha, datos.numero
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    return {
        "mensaje": "Resultado agregado correctamente",
        "vista": servicios_viernes.construir_vista_viernes(db),
    }


@app.post("/viernes/importar-excel")
async def importar_excel_viernes_endpoint(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    if not file.filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        raise HTTPException(status_code=400, detail="El archivo debe ser un Excel válido")

    contenido = await file.read()
    archivo = servicios_viernes.crear_archivo_viernes(db, file.filename)

    try:
        procesadas, diagnostico = servicios_viernes.importar_excel_viernes(db, archivo.id, contenido)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo procesar el archivo: {e}")

    from collections import Counter
    conteo_por_loteria = Counter(loteria for loteria, _ in procesadas)

    return {
        "mensaje": f"{len(procesadas)} filas importadas, {len(diagnostico)} filas omitidas",
        "conteo_por_loteria": dict(conteo_por_loteria),
        "archivo_id": archivo.id,
        "omitidas": diagnostico,
        "vista": servicios_viernes.construir_vista_viernes(db),
    }


@app.get("/viernes/historial/{loteria}")
def obtener_historial_viernes(loteria: str, db: Session = Depends(get_db)):
    resultados = (
        db.query(models_viernes.ResultadoViernes)
        .filter(models_viernes.ResultadoViernes.loteria == loteria)
        .order_by(models_viernes.ResultadoViernes.id)   # <-- orden real del excel
        .all()
    )

    contador_por_numero = {}
    filas = []

    for r in resultados:
        contador_por_numero[r.numero] = contador_por_numero.get(r.numero, 0) + 1
        cantidad = r.cantidad if r.cantidad is not None else contador_por_numero[r.numero]

        filas.append({
            "id": r.id,
            "fecha": r.fecha.isoformat(),
            "numero": r.numero,
            "cantidad": cantidad,
        })

    return filas
@app.post("/historico-4-cifras/sincronizar")
def sincronizar_historico4_endpoint(db: Session = Depends(get_db)):
    """
    Pone al día ResultadoHistorico4 con lo que ya exista en Resultado
    (idempotente: agregar_resultado no duplica nada). Útil como botón
    manual de 'por si acaso' en el frontend, además del sync automático
    que ya corre solo tras la API/scraping diarios.
    """
    agregados = backfill_desde_resultado(db)
    resumen = resumen_historico(db)
    return {
        "mensaje": f"{agregados} resultados nuevos agregados al histórico" if agregados
                   else "El histórico ya estaba al día",
        "resumen": resumen,
    }

@app.get("/viernes/vista")
def obtener_vista_viernes_endpoint(db: Session = Depends(get_db)):
    return servicios_viernes.construir_vista_viernes(db)

@app.get("/debug/scraper-crudo/{nombre_loteria}")
def debug_scraper_crudo(nombre_loteria: str):
    from scrapers.registry import SCRAPERS_ACTIVOS
    scraper = next((s for s in SCRAPERS_ACTIVOS if s.nombre_loteria == nombre_loteria), None)
    if scraper is None:
        return {"error": "no encontrado", "disponibles": [s.nombre_loteria for s in SCRAPERS_ACTIVOS]}

    historico = scraper.obtener_historico(max_sorteos=30)
    return {
        "loteria": nombre_loteria,
        "cantidad": len(historico) if historico else 0,
        "resultados": [
            {"fecha": r.fecha.isoformat(), "numero": r.numero_completo}
            for r in (historico or [])
        ],
    }

@app.get("/permutantes/top")
def top_permutantes(limite: int = 10, db: Session = Depends(get_db)):
    """Los grupos de permutación que más han caído en el histórico,
    con cada número que cayó y su fecha."""
    filas = (
        db.query(
            models_historico4.ResultadoHistorico4.fecha,
            models_historico4.ResultadoHistorico4.numero,
            models_historico4.ResultadoHistorico4.loteria,
        )
        .filter(
            models_historico4.ResultadoHistorico4.numero.isnot(None),
            models_historico4.ResultadoHistorico4.fecha.isnot(None),
        )
        .all()
    )

    def clave_grupo(numero: str) -> str:
        # dígitos de menor a mayor, con el 0 como el mayor
        return "".join(sorted(numero, key=lambda c: 10 if c == "0" else int(c)))

    grupos = {}
    for fecha, numero, loteria in filas:
        if len(numero) != 4 or not numero.isdigit():
            continue  # ignora filas raras (asteriscos, etc.)
        grupos.setdefault(clave_grupo(numero), []).append(
            {"fecha": fecha, "numero": numero, "loteria": loteria}
        )

    ranking = sorted(grupos.items(), key=lambda kv: len(kv[1]), reverse=True)[:limite]

    resultado = []
    for grupo, apariciones in ranking:
        apariciones.sort(key=lambda a: a["fecha"], reverse=True)  # más reciente primero
        resultado.append({
            "grupo": grupo,
            "cantidad": len(apariciones),
            "ultima_fecha": apariciones[0]["fecha"].isoformat(),
            "primera_fecha": apariciones[-1]["fecha"].isoformat(),
            "apariciones": [
                {
                    "fecha": a["fecha"].isoformat(),
                    "numero": a["numero"],
                    "loteria": a["loteria"],
                }
                for a in apariciones
            ],
        })

    return {"top": resultado}

@app.get("/historico-4-cifras/duplicados")
def ver_duplicados_historico4(db: Session = Depends(get_db)):
    """Vista previa: grupos con la misma fecha y el mismo número. No borra nada."""
    H = models_historico4.ResultadoHistorico4
    grupos = (
        db.query(H.fecha, H.numero, func.count(H.id))
        .filter(H.fecha.isnot(None), H.numero.isnot(None))
        .group_by(H.fecha, H.numero)
        .having(func.count(H.id) > 1)
        .order_by(H.fecha.desc())
        .all()
    )
    return {
        "grupos_duplicados": len(grupos),
        "filas_que_se_eliminarian": sum(c - 1 for _, _, c in grupos),
        "detalle": [
            {"fecha": f.isoformat(), "numero": n, "veces": c}
            for f, n, c in grupos
        ],
    }


@app.post("/historico-4-cifras/eliminar-duplicados")
def eliminar_duplicados_historico4(db: Session = Depends(get_db)):
    """Deja UNA fila por cada (fecha, número) y borra las demás.
    Se conserva la de menor id, que es la más antigua (la del Excel)."""
    resultado = db.execute(text("""
        DELETE FROM resultados_historico_4cifras a
        USING resultados_historico_4cifras b
        WHERE a.fecha = b.fecha
          AND a.numero = b.numero
          AND a.id > b.id
    """))
    db.commit()
    return {
        "mensaje": f"{resultado.rowcount} filas duplicadas eliminadas",
        "resumen": resumen_historico(db),
    }    

@app.get("/permutantes/conteos")
def conteos_permutantes(db: Session = Depends(get_db)):
    """Veces que ha caído cada número de 4 cifras y la última fecha en que cayó."""
    H = models_historico4.ResultadoHistorico4
    filas = (
        db.query(H.numero, func.count(H.id), func.max(H.fecha))
        .filter(H.numero.isnot(None))
        .group_by(H.numero)
        .all()
    )
    # Ignora filas raras (con asterisco, sin 4 cifras, etc.)
    validas = [(n, c, f) for n, c, f in filas if len(n) == 4 and n.isdigit()]
    return {
        "conteos": {n: c for n, c, _ in validas},
        "ultimas_fechas": {n: f.isoformat() for n, _, f in validas if f},
        "total": sum(c for _, c, _ in validas),
    }

 