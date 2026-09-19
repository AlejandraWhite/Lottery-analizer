"""
services/historico4.py

Histórico de 4 cifras (Excel de ~20 años). Sin lógica de análisis:
solo se guarda lo que trae el Excel (columnas B..J) y luego se van
agregando los números nuevos que caen.

Regla de oro del import: NO se omite nada. Si una fila trae algo raro
se guarda igual y se anota en `observacion`.
"""

import re
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import func
from sqlalchemy.orm import Session

from models_historico4 import ResultadoHistorico4

MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9,
    "octubre": 10, "noviembre": 11, "diciembre": 12,
}

# Columnas D..J (las que van después de B=fecha y C=número)
LETRAS_EXTRA = ["D", "E", "F", "G", "H", "I", "J"]


class HistoricoYaImportado(Exception):
    pass


# =========================================================
# PARSEO
# =========================================================

# Todos capturan al final un "sobrante": texto pegado a la fecha
# (ej. "04/09/2023Ext Manizal").
_P_MES_DIA_ANIO = re.compile(r"^([a-záéíóúñ]+)\.?\s*(\d{1,2})\s*,?\s*(\d{4})(.*)$", re.IGNORECASE)
_P_DIA_DE_MES = re.compile(r"^(\d{1,2})\s+de\s+([a-záéíóúñ]+)\s+de\s+(\d{4})(.*)$", re.IGNORECASE)
_P_DMA = re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})(.*)$")
_P_ISO = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(.*)$")
_P_HORA = re.compile(r"^[ tT]\d{2}:\d{2}(:\d{2})?(\.\d+)?$")


def extraer_fecha(valor) -> tuple[date, str]:
    """Devuelve (fecha, texto_sobrante). Acepta datetime/date de Excel,
    "julio 05, 2019", "mayo 17,2019", "8/07/2019" (dd/mm/yyyy),
    "2019-07-08" y "05 de julio de 2019", con o sin texto pegado detrás."""
    if isinstance(valor, datetime):
        return valor.date(), ""
    if isinstance(valor, date):
        return valor, ""
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        if 20000 < valor < 80000:  # número de serie de Excel
            return from_excel(valor).date(), ""
        raise ValueError(f"fecha numérica fuera de rango: {valor!r}")

    texto = str(valor).strip()

    m = _P_MES_DIA_ANIO.match(texto)
    if m and m.group(1).lower() in MESES:
        return date(int(m.group(3)), MESES[m.group(1).lower()], int(m.group(2))), m.group(4).strip()

    m = _P_DIA_DE_MES.match(texto)
    if m and m.group(2).lower() in MESES:
        return date(int(m.group(3)), MESES[m.group(2).lower()], int(m.group(1))), m.group(4).strip()

    m = _P_DMA.match(texto)
    if m:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1))), m.group(4).strip()

    m = _P_ISO.match(texto)
    if m:
        resto = m.group(4)
        if _P_HORA.match(resto):
            resto = ""
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3))), resto.strip()

    raise ValueError(f"formato de fecha no reconocido: {valor!r}")


def parsear_fecha(valor) -> date:
    """Versión estricta: falla si hay texto pegado a la fecha."""
    fecha, sobrante = extraer_fecha(valor)
    if sobrante:
        raise ValueError(f"texto extra junto a la fecha: {sobrante!r}")
    return fecha


def normalizar_numero(valor) -> str:
    """Devuelve el número como texto de 4 cifras, recuperando ceros
    a la izquierda que Excel haya perdido (883 -> "0883")."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    if not texto.isdigit() or len(texto) > 4:
        raise ValueError(f"no es un número de 4 cifras: {valor!r}")
    return texto.zfill(4)


def _valor_json(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, (int, float, str)):
        return v
    return str(v)


def _extras(celdas) -> dict | None:
    extras = {}
    for letra, v in zip(LETRAS_EXTRA, celdas):
        if v is None or str(v).strip() == "":
            continue
        extras[letra] = _valor_json(v)
    return extras or None


def _hay(valor) -> bool:
    return valor is not None and str(valor).strip() != ""


# =========================================================
# IMPORTAR EXCEL
# =========================================================

def importar_historico(
    db: Session,
    contenido: bytes,
    nombre_archivo: str,
    reemplazar: bool = False,
    fila_inicial: int = 1,
):
    """
    Lee columnas B..J de la primera hoja.
      B = fecha (si está vacía, se usa la última fecha vista hacia arriba:
          varios sorteos del mismo día comparten la fecha de la primera fila)
      C = número de 4 cifras
      D..J = se guardan tal cual en columnas_extra (sin ninguna lógica)

    NO se omite ninguna fila con contenido. Lo raro se guarda igual y
    queda anotado en `observacion`:
      - texto pegado a la fecha  -> fecha buena + nota con el texto
      - fecha que no se entiende -> fecha NULL + nota con el original
      - número con caracteres raros -> se guarda tal cual + nota
      - fila sin número en C     -> numero NULL + nota
    Solo se saltan las filas completamente vacías en B..J.

    Devuelve (cantidad_insertada, con_observaciones).
    Si ya hay histórico cargado y reemplazar=False, lanza HistoricoYaImportado
    para no duplicar 20 años de datos por accidente.
    """
    existentes = db.query(func.count(ResultadoHistorico4.id)).scalar()
    if existentes and not reemplazar:
        raise HistoricoYaImportado(
            f"Ya hay {existentes} filas en el histórico. "
            "Usa reemplazar=true si quieres borrarlo y cargarlo de nuevo."
        )

    workbook = load_workbook(filename=BytesIO(contenido), read_only=True, data_only=True)
    try:
        hoja = workbook.worksheets[0]

        filas = []
        con_observaciones = []
        ultima_fecha = None

        for n, celdas in enumerate(
            hoja.iter_rows(min_row=fila_inicial, min_col=2, max_col=10, values_only=True),
            start=fila_inicial,
        ):
            celdas = list(celdas) + [None] * (9 - len(celdas))
            valor_fecha, valor_numero, *resto = celdas

            hay_fecha = _hay(valor_fecha)
            hay_numero = _hay(valor_numero)
            extras = _extras(resto)

            if not hay_fecha and not hay_numero and not extras:
                continue  # fila realmente vacía

            notas = []

            # --- Fecha ---
            fecha = None
            if hay_fecha:
                try:
                    fecha, sobrante = extraer_fecha(valor_fecha)
                    if sobrante:
                        notas.append(f"Texto pegado a la fecha en el Excel: '{valor_fecha}'")
                    ultima_fecha = fecha
                except Exception:
                    ultima_fecha = None  # no arrastrar una fecha dudosa a las filas de abajo
                    notas.append(f"Fecha no reconocida en el Excel: '{valor_fecha}'")
            elif hay_numero:
                if ultima_fecha is not None:
                    fecha = ultima_fecha  # mismo día que la fila de arriba
                else:
                    notas.append("Sin fecha en el Excel y sin fecha anterior para heredar")

            # --- Número ---
            numero = None
            if hay_numero:
                try:
                    numero = normalizar_numero(valor_numero)
                except ValueError:
                    numero = str(valor_numero).strip()[:20]
                    notas.append(f"Número no estándar en el Excel: '{valor_numero}'")
            else:
                notas.append("Fila sin número en la columna C")

            observacion = "; ".join(notas) if notas else None

            filas.append({
                "fila_excel": n,
                "nombre_archivo": nombre_archivo,
                "fecha": fecha,
                "numero": numero,
                "loteria": None,
                "columnas_extra": extras,
                "observacion": observacion,
            })
            if observacion:
                con_observaciones.append({
                    "fila_excel": n,
                    "fecha": str(valor_fecha) if hay_fecha else None,
                    "numero": str(valor_numero) if hay_numero else None,
                    "observacion": observacion,
                })

        if reemplazar:
            db.query(ResultadoHistorico4).delete()
        db.bulk_insert_mappings(ResultadoHistorico4, filas)
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        workbook.close()

    return len(filas), con_observaciones


# =========================================================
# AGREGAR NÚMEROS NUEVOS (manual / scraping / API)
# =========================================================

def agregar_resultado(
    db: Session,
    fecha,
    numero_completo,
    loteria: str | None = None,
    commit: bool = True,
):
    """Agrega un número nuevo al histórico. Idempotente: si ya existe
    (misma fecha + número) no lo duplica. Ignora en silencio lo que no
    sea de 4 cifras."""
    if numero_completo is None or len(str(numero_completo).strip()) != 4:
        return None
    try:
        numero = normalizar_numero(numero_completo)
        if isinstance(fecha, datetime):
            fecha = fecha.date()
        elif not isinstance(fecha, date):
            fecha = parsear_fecha(fecha)
    except ValueError:
        return None

    # Se deduplica por fecha + número (sin lotería), igual que
    # existe_resultado_global: así no se duplica lo que ya venga en el Excel
    # histórico (que no trae lotería) ni lo que llegue por dos vías.
    ya_existe = (
        db.query(ResultadoHistorico4.id)
        .filter(ResultadoHistorico4.fecha == fecha, ResultadoHistorico4.numero == numero)
        .first()
    )
    if ya_existe:
        return None

    fila = ResultadoHistorico4(fecha=fecha, numero=numero, loteria=loteria)
    db.add(fila)
    if commit:
        db.commit()
    else:
        db.flush()
    return fila


def agregar_desde_resultado(db: Session, resultado, commit: bool = True):
    """Atajo: recibe un objeto Resultado (el de models.py)."""
    return agregar_resultado(
        db, resultado.fecha, resultado.numero_completo, resultado.loteria, commit=commit
    )


def quitar_resultado(db: Session, fecha, numero_completo, commit: bool = False) -> int:
    """Quita del histórico un número que se agregó automáticamente (se usa al
    eliminar un resultado). Nunca toca filas que vinieron del Excel
    (fila_excel no nulo). Devuelve cuántas filas borró."""
    try:
        numero = normalizar_numero(numero_completo)
        if isinstance(fecha, datetime):
            fecha = fecha.date()
    except ValueError:
        return 0

    borradas = (
        db.query(ResultadoHistorico4)
        .filter(
            ResultadoHistorico4.fecha == fecha,
            ResultadoHistorico4.numero == numero,
            ResultadoHistorico4.fila_excel.is_(None),
        )
        .delete(synchronize_session=False)
    )
    if commit:
        db.commit()
    return borradas


def resumen_historico(db: Session) -> dict:
    total, primera, ultima = db.query(
        func.count(ResultadoHistorico4.id),
        func.min(ResultadoHistorico4.fecha),
        func.max(ResultadoHistorico4.fecha),
    ).one()
    return {
        "total": total,
        "primera_fecha": primera.isoformat() if primera else None,
        "ultima_fecha": ultima.isoformat() if ultima else None,
    }

# =========================================================
# BACKFILL (poner al día ResultadoHistorico4 con lo que ya
# existe en la tabla Resultado, por si entró antes de que
# el sync/scraping empezara a alimentar el histórico)
# =========================================================

def backfill_desde_resultado(db: Session) -> int:
    """
    Recorre TODA la tabla Resultado y agrega a ResultadoHistorico4 lo
    que falte (agregar_resultado ya es idempotente por fecha+numero,
    así que no duplica nada). Útil una sola vez después de conectar el
    sync/scraping al histórico, para que se ponga al día con lo que ya
    había entrado antes del cambio.
    """
    from models import Resultado  # import local para evitar ciclo de imports

    resultados = db.query(Resultado).order_by(Resultado.fecha.asc()).all()
    agregados = 0

    for r in resultados:
        fila = agregar_resultado(db, r.fecha, r.numero_completo, r.loteria, commit=False)
        if fila is not None:
            agregados += 1

    db.commit()
    return agregados