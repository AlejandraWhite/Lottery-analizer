from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from models_miercoles import (
    ArchivoMiercoles,
    ResultadoMiercoles,
    ContadorTerminacionMiercoles,
    EstadoTerminacionMiercoles,
)
from services.analisis import normalizar_numero

FILA_INICIO = 2
MAX_FILAS = 500  # tope de seguridad, no un requisito de exactamente 100


# =========================================================
# FECHA: parser tolerante
# =========================================================

def _parsear_fecha_flexible(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor

    if isinstance(valor, (int, float)):
        dias = round(valor)  # evita el desfase de 1 día por precisión de float
        return (datetime(1899, 12, 30) + timedelta(days=dias)).date()

    texto = str(valor).strip()
    if not texto:
        raise ValueError("Fecha vacía")

    formatos = (
        "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y",
        "%d/%m/%y", "%d-%m-%y",
    )
    for formato in formatos:
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue

    raise ValueError(f"No se pudo interpretar la fecha: {valor!r}")


# =========================================================
# ARCHIVO
# =========================================================

def crear_archivo_miercoles(db: Session, nombre: str) -> ArchivoMiercoles:
    ahora = datetime.now()
    archivo = ArchivoMiercoles(
        nombre=nombre,
        nombre_original=nombre,
        fecha_creacion=ahora,
        fecha_actualizacion=ahora,
    )
    db.add(archivo)
    db.commit()
    db.refresh(archivo)
    return archivo


NOMBRE_ARCHIVO_MANUAL_MIERCOLES = "Ingreso manual (miércoles)"


def obtener_o_crear_archivo_manual_miercoles(db: Session) -> ArchivoMiercoles:
    archivo = (
        db.query(ArchivoMiercoles)
        .filter(ArchivoMiercoles.nombre == NOMBRE_ARCHIVO_MANUAL_MIERCOLES)
        .first()
    )
    if archivo:
        return archivo
    return crear_archivo_miercoles(db, NOMBRE_ARCHIVO_MANUAL_MIERCOLES)


# =========================================================
# ESTADO / CONTADOR POR (LOTERIA, TERMINACION)
# =========================================================

def obtener_o_crear_estado(db: Session, loteria: str, terminacion: str) -> EstadoTerminacionMiercoles:
    estado = (
        db.query(EstadoTerminacionMiercoles)
        .filter(
            EstadoTerminacionMiercoles.loteria == loteria,
            EstadoTerminacionMiercoles.terminacion == terminacion,
        )
        .first()
    )
    if estado is None:
        estado = EstadoTerminacionMiercoles(loteria=loteria, terminacion=terminacion)
        db.add(estado)
        db.flush()
    return estado


def obtener_o_crear_contador(db: Session, loteria: str, terminacion: str) -> ContadorTerminacionMiercoles:
    contador = (
        db.query(ContadorTerminacionMiercoles)
        .filter(
            ContadorTerminacionMiercoles.loteria == loteria,
            ContadorTerminacionMiercoles.terminacion == terminacion,
        )
        .first()
    )
    if contador is None:
        contador = ContadorTerminacionMiercoles(loteria=loteria, terminacion=terminacion, cantidad=0)
        db.add(contador)
        db.flush()
    return contador


def _obtener_o_crear_resultado(db: Session, archivo_id: int, loteria: str, fecha, numero: str) -> ResultadoMiercoles:
    resultado = (
        db.query(ResultadoMiercoles)
        .filter(
            ResultadoMiercoles.loteria == loteria,
            ResultadoMiercoles.fecha == fecha,
            ResultadoMiercoles.numero == numero,
        )
        .first()
    )
    if resultado is None:
        resultado = ResultadoMiercoles(
            archivo_id=archivo_id, fecha=fecha, numero=numero, loteria=loteria
        )
        db.add(resultado)
        db.flush()
    return resultado


def _eliminar_si_existe(db: Session, resultado_id):
    if resultado_id is None:
        return
    vieja = db.get(ResultadoMiercoles, resultado_id)
    if vieja is not None:
        db.delete(vieja)
        db.flush()


# =========================================================
# CADENA GENÉRICA (Valle / Manizales / ingreso manual)
# última -> penúltima -> antepenúltima -> se elimina
# =========================================================

def _registrar_aparicion(db: Session, archivo_id: int, loteria: str, fecha, numero: str,
                          cantidad_forzada=None) -> ResultadoMiercoles:
    estado = obtener_o_crear_estado(db, loteria, numero)

    if estado.ultima_id is not None:
        actual = db.get(ResultadoMiercoles, estado.ultima_id)
        if actual is not None and actual.fecha == fecha and actual.numero == numero:
            return actual  # ya está registrada esta misma aparición, no hay nada nuevo

    id_antepenultima_anterior = estado.antepenultima_id
    id_penultima_anterior = estado.penultima_id
    id_ultima_anterior = estado.ultima_id

    nuevo = _obtener_o_crear_resultado(db, archivo_id, loteria, fecha, numero)

    estado.antepenultima_id = id_penultima_anterior
    estado.penultima_id = id_ultima_anterior
    estado.ultima_id = nuevo.id
    db.flush()

    _eliminar_si_existe(db, id_antepenultima_anterior)

    contador = obtener_o_crear_contador(db, loteria, numero)
    contador.cantidad = cantidad_forzada if cantidad_forzada is not None else contador.cantidad + 1
    nuevo.cantidad = contador.cantidad

    return nuevo


def registrar_resultado_manual_miercoles(db: Session, loteria: str, fecha_texto, numero) -> ResultadoMiercoles:
    if loteria not in ("Meta", "Valle", "Manizales"):
        raise ValueError("Lotería inválida: debe ser Meta, Valle o Manizales")

    fecha = _parsear_fecha_flexible(fecha_texto)
    terminacion = normalizar_numero(numero)

    archivo = obtener_o_crear_archivo_manual_miercoles(db)
    nuevo = _registrar_aparicion(db, archivo.id, loteria, fecha, terminacion)

    db.commit()
    db.refresh(nuevo)
    return nuevo


# =========================================================
# META: última y penúltima SIEMPRE se leen de K/L y N/O
# (fuente de verdad). Antepenúltima se calcula sola:
# cuando cambia la última, la penúltima que había pasa a
# antepenúltima, y la antepenúltima vieja se elimina.
# =========================================================

def _procesar_fila_meta(db: Session, archivo_id: int, fila: dict):
    numero = fila["numero"]
    fecha_nueva = fila["fecha"]

    estado = obtener_o_crear_estado(db, "Meta", numero)
    contador = obtener_o_crear_contador(db, "Meta", numero)

    ultima_actual = db.get(ResultadoMiercoles, estado.ultima_id) if estado.ultima_id else None
    es_sorteo_nuevo = ultima_actual is None or ultima_actual.fecha != fecha_nueva

    if es_sorteo_nuevo:
        # lo que estaba en penúltima se corre a antepenúltima;
        # lo que había antes en antepenúltima se elimina
        _eliminar_si_existe(db, estado.antepenultima_id)
        estado.antepenultima_id = estado.penultima_id

    # la penúltima siempre se confía directo del excel (N/O)
    if fila["fecha_pen"] is not None and fila["numero_pen"] is not None:
        resultado_pen = _obtener_o_crear_resultado(
            db, archivo_id, "Meta", fila["fecha_pen"], fila["numero_pen"]
        )
        estado.penultima_id = resultado_pen.id
    elif es_sorteo_nuevo:
        # el excel no trajo N/O para esta fila: al menos conservamos
        # la que era la última como penúltima
        estado.penultima_id = ultima_actual.id if ultima_actual else None

    nueva_ultima = _obtener_o_crear_resultado(db, archivo_id, "Meta", fecha_nueva, numero)
    estado.ultima_id = nueva_ultima.id

    if fila["cantidad"] is not None:
        contador.cantidad = fila["cantidad"]
    elif es_sorteo_nuevo:
        contador.cantidad += 1

    nueva_ultima.cantidad = contador.cantidad


# =========================================================
# LECTURA DE FILAS POR LOTERIA (con diagnóstico de omitidas)
# =========================================================

def _valor(ws, columna, fila):
    return ws[f"{columna}{fila}"].value


def _celda_vacia(valor):
    return valor is None or str(valor).strip() == ""


def _leer_filas_meta(ws, diagnostico, fila_inicio=FILA_INICIO, max_filas=MAX_FILAS):
    filas = []
    for i in range(max_filas):
        fila = fila_inicio + i
        celda_fecha = _valor(ws, "K", fila)
        celda_numero = _valor(ws, "L", fila)
        celda_cantidad = _valor(ws, "M", fila)
        celda_fecha_pen = _valor(ws, "N", fila)
        celda_numero_pen = _valor(ws, "O", fila)

        if _celda_vacia(celda_fecha) and _celda_vacia(celda_numero):
            continue

        if _celda_vacia(celda_fecha) or _celda_vacia(celda_numero):
            diagnostico.append(f"Meta fila {fila}: K={celda_fecha!r} L={celda_numero!r}")
            continue

        try:
            fecha = _parsear_fecha_flexible(celda_fecha)
            numero = normalizar_numero(celda_numero)
        except (ValueError, TypeError) as e:
            diagnostico.append(f"Meta fila {fila}: K={celda_fecha!r} L={celda_numero!r} -> {e}")
            continue

        cantidad = None
        if not _celda_vacia(celda_cantidad):
            try:
                cantidad = int(celda_cantidad)
            except (ValueError, TypeError):
                cantidad = None

        fecha_pen = numero_pen = None
        if not _celda_vacia(celda_fecha_pen) and not _celda_vacia(celda_numero_pen):
            try:
                fecha_pen = _parsear_fecha_flexible(celda_fecha_pen)
                numero_pen = normalizar_numero(celda_numero_pen)
            except (ValueError, TypeError):
                fecha_pen = numero_pen = None

        filas.append({
            "fecha": fecha,
            "numero": numero,
            "cantidad": cantidad,
            "fecha_pen": fecha_pen,
            "numero_pen": numero_pen,
        })
    return filas


def _leer_filas_simple(ws, col_fecha, col_numero, nombre_loteria, diagnostico,
                        fila_inicio=FILA_INICIO, max_filas=MAX_FILAS):
    filas = []
    for i in range(max_filas):
        fila = fila_inicio + i
        celda_fecha = _valor(ws, col_fecha, fila)
        celda_numero = _valor(ws, col_numero, fila)

        if _celda_vacia(celda_fecha) and _celda_vacia(celda_numero):
            continue

        if _celda_vacia(celda_fecha) or _celda_vacia(celda_numero):
            diagnostico.append(
                f"{nombre_loteria} fila {fila}: {col_fecha}={celda_fecha!r} {col_numero}={celda_numero!r}"
            )
            continue

        try:
            fecha = _parsear_fecha_flexible(celda_fecha)
            numero = normalizar_numero(celda_numero)
        except (ValueError, TypeError) as e:
            diagnostico.append(
                f"{nombre_loteria} fila {fila}: {col_fecha}={celda_fecha!r} {col_numero}={celda_numero!r} -> {e}"
            )
            continue

        filas.append({"fecha": fecha, "numero": numero})
    return filas


# =========================================================
# IMPORTAR DESDE EXCEL
# =========================================================

def importar_excel_miercoles(db: Session, archivo_id: int, contenido_excel: bytes):
    wb = load_workbook(filename=BytesIO(contenido_excel), data_only=True)
    ws = wb.active

    procesadas = []
    diagnostico = []

    # ---- META: K/L = última, N/O = penúltima (siempre confiadas del excel) ----
    for fila in _leer_filas_meta(ws, diagnostico):
        _procesar_fila_meta(db, archivo_id, fila)
        procesadas.append(("Meta", fila["numero"]))

    # ---- VALLE y MANIZALES: cadena genérica (el excel no trae penúltima) ----
    for loteria, col_fecha, col_numero in (
        ("Valle", "Q", "R"),
        ("Manizales", "T", "U"),
    ):
        for fila in _leer_filas_simple(ws, col_fecha, col_numero, loteria, diagnostico):
            _registrar_aparicion(db, archivo_id, loteria, fila["fecha"], fila["numero"])
            procesadas.append((loteria, fila["numero"]))

    db.commit()
    return procesadas, diagnostico


# =========================================================
# VISTA PARA EL FRONTEND (una fila por terminación, 00-99)
# =========================================================

def _convertir(resultado: ResultadoMiercoles):
    if resultado is None:
        return None
    return {
        "id": resultado.id,
        "fecha": resultado.fecha.isoformat() if resultado.fecha else None,
        "numero": resultado.numero,
    }


def construir_vista_miercoles(db: Session):
    vista = {}

    for loteria in ("Meta", "Valle", "Manizales"):
        estados = {
            e.terminacion: e
            for e in db.query(EstadoTerminacionMiercoles)
            .filter(EstadoTerminacionMiercoles.loteria == loteria)
            .all()
        }
        contadores = {
            c.terminacion: c.cantidad
            for c in db.query(ContadorTerminacionMiercoles)
            .filter(ContadorTerminacionMiercoles.loteria == loteria)
            .all()
        }

        filas = []
        for i in range(100):
            terminacion = f"{i:02d}"
            estado = estados.get(terminacion)

            filas.append({
                "terminacion": terminacion,
                "cantidad": contadores.get(terminacion, 0),
                "ultima": _convertir(estado.ultima) if estado else None,
                "penultima": _convertir(estado.penultima) if estado else None,
                "antepenultima": _convertir(estado.antepenultima) if estado else None,
            })

        filas.sort(
            key=lambda f: f["ultima"]["fecha"] if f["ultima"] else "9999-99-99"
        )
        vista[loteria] = filas

    return vista