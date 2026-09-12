from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from models_viernes import (
    ArchivoViernes,
    ResultadoViernes,
    ContadorTerminacionViernes,
    EstadoTerminacionViernes,
)
from services.analisis import normalizar_numero

FILA_INICIO = 2
MAX_FILAS = 500  # tope de seguridad, no un requisito de exactamente 100

LOTERIAS_VIERNES = ("Risaralda", "Medellin", "Santander")

# columnas del excel por loteria: (fecha, numero, cantidad)
COLUMNAS_VIERNES = {
    "Risaralda": ("L", "M", "N"),
    "Medellin": ("R", "S", "T"),
    "Santander": ("X", "Y", "Z"),
}


# =========================================================
# FECHA: parser tolerante (idéntico al de miércoles)
# =========================================================

def _parsear_fecha_flexible(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor

    if isinstance(valor, (int, float)):
        dias = round(valor)  # evita el desfase de 1 día por precisión de float
        return (datetime(1899, 12, 30) + timedelta(days=dias)).date()

    texto = str(valor).strip()
    if not texto:
        raise ValueError("Fecha vacía")

    formatos = (
        "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y",
        "%d/%m/%y", "%d-%m-%y",
    )
    for formato in formatos:
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue

    raise ValueError(f"No se pudo interpretar la fecha: {valor!r}")


# =========================================================
# ARCHIVO
# =========================================================

def crear_archivo_viernes(db: Session, nombre: str) -> ArchivoViernes:
    ahora = datetime.now()
    archivo = ArchivoViernes(
        nombre=nombre,
        nombre_original=nombre,
        fecha_creacion=ahora,
        fecha_actualizacion=ahora,
    )
    db.add(archivo)
    db.commit()
    db.refresh(archivo)
    return archivo


NOMBRE_ARCHIVO_MANUAL_VIERNES = "Ingreso manual (viernes)"


def obtener_o_crear_archivo_manual_viernes(db: Session) -> ArchivoViernes:
    archivo = (
        db.query(ArchivoViernes)
        .filter(ArchivoViernes.nombre == NOMBRE_ARCHIVO_MANUAL_VIERNES)
        .first()
    )
    if archivo:
        return archivo
    return crear_archivo_viernes(db, NOMBRE_ARCHIVO_MANUAL_VIERNES)


# =========================================================
# ESTADO / CONTADOR POR (LOTERIA, TERMINACION)
# =========================================================

def obtener_o_crear_estado(db: Session, loteria: str, terminacion: str) -> EstadoTerminacionViernes:
    estado = (
        db.query(EstadoTerminacionViernes)
        .filter(
            EstadoTerminacionViernes.loteria == loteria,
            EstadoTerminacionViernes.terminacion == terminacion,
        )
        .first()
    )
    if estado is None:
        estado = EstadoTerminacionViernes(loteria=loteria, terminacion=terminacion)
        db.add(estado)
        db.flush()
    return estado


def obtener_o_crear_contador(db: Session, loteria: str, terminacion: str) -> ContadorTerminacionViernes:
    contador = (
        db.query(ContadorTerminacionViernes)
        .filter(
            ContadorTerminacionViernes.loteria == loteria,
            ContadorTerminacionViernes.terminacion == terminacion,
        )
        .first()
    )
    if contador is None:
        contador = ContadorTerminacionViernes(loteria=loteria, terminacion=terminacion, cantidad=0)
        db.add(contador)
        db.flush()
    return contador


def _obtener_o_crear_resultado(db: Session, archivo_id: int, loteria: str, fecha, numero: str) -> ResultadoViernes:
    resultado = (
        db.query(ResultadoViernes)
        .filter(
            ResultadoViernes.loteria == loteria,
            ResultadoViernes.fecha == fecha,
            ResultadoViernes.numero == numero,
        )
        .first()
    )
    if resultado is None:
        resultado = ResultadoViernes(
            archivo_id=archivo_id, fecha=fecha, numero=numero, loteria=loteria
        )
        db.add(resultado)
        db.flush()
    return resultado


def _eliminar_si_existe(db: Session, resultado_id):
    if resultado_id is None:
        return
    vieja = db.get(ResultadoViernes, resultado_id)
    if vieja is not None:
        db.delete(vieja)
        db.flush()


# =========================================================
# CADENA GENÉRICA (Risaralda / Medellín / Santander / ingreso manual)
# última -> penúltima -> antepenúltima -> se elimina
#
# Las 3 loterías del viernes traen su propia columna de "cantidad"
# en el excel (N/T/Z), así que si viene un valor se usa como
# cantidad_forzada; si no, se sigue incrementando el contador.
# =========================================================

def _registrar_aparicion(db: Session, archivo_id: int, loteria: str, fecha, numero: str,
                          cantidad_forzada=None) -> ResultadoViernes:
    estado = obtener_o_crear_estado(db, loteria, numero)

    if estado.ultima_id is not None:
        actual = db.get(ResultadoViernes, estado.ultima_id)
        if actual is not None and actual.fecha == fecha and actual.numero == numero:
            return actual  # ya está registrada esta misma aparición, no hay nada nuevo

    id_antepenultima_anterior = estado.antepenultima_id
    id_penultima_anterior = estado.penultima_id
    id_ultima_anterior = estado.ultima_id

    nuevo = _obtener_o_crear_resultado(db, archivo_id, loteria, fecha, numero)

    estado.antepenultima_id = id_penultima_anterior
    estado.penultima_id = id_ultima_anterior
    estado.ultima_id = nuevo.id
    db.flush()

    _eliminar_si_existe(db, id_antepenultima_anterior)

    contador = obtener_o_crear_contador(db, loteria, numero)
    contador.cantidad = cantidad_forzada if cantidad_forzada is not None else contador.cantidad + 1
    nuevo.cantidad = contador.cantidad

    return nuevo


def registrar_resultado_manual_viernes(db: Session, loteria: str, fecha_texto, numero) -> ResultadoViernes:
    if loteria not in LOTERIAS_VIERNES:
        raise ValueError("Lotería inválida: debe ser Risaralda, Medellin o Santander")

    fecha = _parsear_fecha_flexible(fecha_texto)
    terminacion = normalizar_numero(numero)

    archivo = obtener_o_crear_archivo_manual_viernes(db)
    nuevo = _registrar_aparicion(db, archivo.id, loteria, fecha, terminacion)

    db.commit()
    db.refresh(nuevo)
    return nuevo


# =========================================================
# LECTURA DE FILAS POR LOTERIA (con diagnóstico de omitidas)
# =========================================================

def _valor(ws, columna, fila):
    return ws[f"{columna}{fila}"].value


def _celda_vacia(valor):
    return valor is None or str(valor).strip() == ""


def _leer_filas_loteria(ws, loteria, diagnostico, fila_inicio=FILA_INICIO, max_filas=MAX_FILAS):
    col_fecha, col_numero, col_cantidad = COLUMNAS_VIERNES[loteria]
    filas = []
    for i in range(max_filas):
        fila = fila_inicio + i
        celda_fecha = _valor(ws, col_fecha, fila)
        celda_numero = _valor(ws, col_numero, fila)
        celda_cantidad = _valor(ws, col_cantidad, fila)

        if _celda_vacia(celda_fecha) and _celda_vacia(celda_numero):
            continue

        if _celda_vacia(celda_fecha) or _celda_vacia(celda_numero):
            diagnostico.append(
                f"{loteria} fila {fila}: {col_fecha}={celda_fecha!r} {col_numero}={celda_numero!r}"
            )
            continue

        try:
            fecha = _parsear_fecha_flexible(celda_fecha)
            numero = normalizar_numero(celda_numero)
        except (ValueError, TypeError) as e:
            diagnostico.append(
                f"{loteria} fila {fila}: {col_fecha}={celda_fecha!r} {col_numero}={celda_numero!r} -> {e}"
            )
            continue

        cantidad = None
        if not _celda_vacia(celda_cantidad):
            try:
                cantidad = int(celda_cantidad)
            except (ValueError, TypeError):
                cantidad = None

        filas.append({"fecha": fecha, "numero": numero, "cantidad": cantidad})
    return filas


# =========================================================
# IMPORTAR DESDE EXCEL
# =========================================================

def importar_excel_viernes(db: Session, archivo_id: int, contenido_excel: bytes):
    wb = load_workbook(filename=BytesIO(contenido_excel), data_only=True)
    ws = wb.active

    procesadas = []
    diagnostico = []

    for loteria in LOTERIAS_VIERNES:
        for fila in _leer_filas_loteria(ws, loteria, diagnostico):
            _registrar_aparicion(
                db, archivo_id, loteria, fila["fecha"], fila["numero"],
                cantidad_forzada=fila["cantidad"],
            )
            procesadas.append((loteria, fila["numero"]))

    db.commit()
    return procesadas, diagnostico


# =========================================================
# VISTA PARA EL FRONTEND (una fila por terminación, 00-99)
# =========================================================

def _convertir(resultado: ResultadoViernes):
    if resultado is None:
        return None
    return {
        "id": resultado.id,
        "fecha": resultado.fecha.isoformat() if resultado.fecha else None,
        "numero": resultado.numero,
    }


def construir_vista_viernes(db: Session):
    vista = {}

    for loteria in LOTERIAS_VIERNES:
        estados = {
            e.terminacion: e
            for e in db.query(EstadoTerminacionViernes)
            .filter(EstadoTerminacionViernes.loteria == loteria)
            .all()
        }
        contadores = {
            c.terminacion: c.cantidad
            for c in db.query(ContadorTerminacionViernes)
            .filter(ContadorTerminacionViernes.loteria == loteria)
            .all()
        }

        filas = []
        for i in range(100):
            terminacion = f"{i:02d}"
            estado = estados.get(terminacion)

            filas.append({
                "terminacion": terminacion,
                "cantidad": contadores.get(terminacion, 0),
                "ultima": _convertir(estado.ultima) if estado else None,
                "penultima": _convertir(estado.penultima) if estado else None,
                "antepenultima": _convertir(estado.antepenultima) if estado else None,
            })

        filas.sort(
            key=lambda f: f["ultima"]["fecha"] if f["ultima"] else "9999-99-99"
        )
        vista[loteria] = filas

    return vista


NOMBRE_ARCHIVO_SINCRONIZACION_VIERNES = "Sincronización automática (scraping)"


def obtener_o_crear_archivo_sincronizacion_viernes(db: Session) -> ArchivoViernes:
    archivo = (
        db.query(ArchivoViernes)
        .filter(ArchivoViernes.nombre == NOMBRE_ARCHIVO_SINCRONIZACION_VIERNES)
        .first()
    )
    if archivo:
        return archivo
    return crear_archivo_viernes(db, NOMBRE_ARCHIVO_SINCRONIZACION_VIERNES)


def registrar_resultado_scraping_viernes(db: Session, loteria: str, fecha, numero: str):
    """
    loteria debe ser el nombre corto ("Risaralda" | "Medellin" | "Santander");
    numero, las últimas 2 cifras.
    """
    if loteria not in LOTERIAS_VIERNES:
        return None
    archivo = obtener_o_crear_archivo_sincronizacion_viernes(db)
    return _registrar_aparicion(db, archivo.id, loteria, fecha, numero)